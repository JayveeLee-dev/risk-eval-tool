# -*- coding: utf-8 -*-
"""
评测脚本 v2：调用已部署的工作流接口，直读结构化评分字段，与专家评分对比。
相比 v1 的改进：
  1. 评分直接读取 JSON 结构化字段，不再用正则从报告文本里捞；
  2. 支持中断重跑——已完成的案例自动跳过，不重复调用接口；
  3. 新增"红线预警"正则提取和三个节点分的一致率统计；
  4. 输出"评测结果_v2.xlsx"，不与 v1 的结果文件冲突。
"""

import json
import re
import time

import requests
from openpyxl import load_workbook, Workbook

# ======================= 配置区（按需修改） =======================
API_URL = "https://your-workflow-url.example.com/run"
API_TOKEN = "XXXXXXX"  # ← 把你完整的 Token 粘贴到这里（引号内）
INPUT_FILE = "demo_input.xlsx"
OUTPUT_FILE = "评测结果_v2.xlsx"
INPUT_SHEET = "案例字段"       # 输入数据所在sheet
SUMMARY_SHEET = "评测总表"     # 专家评分所在sheet
REQUEST_TIMEOUT = 90           # 单次请求超时（秒）
ROW_INTERVAL = 3               # 每行请求之间的间隔（秒）
MAX_RETRY = 2                  # 失败后额外重试次数
# ==================================================================

# 12个发给接口的字段（第14列"脱敏检查"不发送）
FIELD_KEYS = [
    "customer_name", "project_name", "contract_amount", "payment_term_days",
    "cooperation_history", "fund_source", "direct_contract", "staff_authority",
    "link_relation", "qualification", "discredit_record", "contract_signer",
]

# 数字字段，读取时做类型转换
NUMERIC_FIELDS = {"contract_amount", "payment_term_days"}

# 输出文件的列（顺序即Excel列顺序）
OUTPUT_HEADER = (["案例ID"] + FIELD_KEYS
                 + ["项目分", "企业分", "合同分", "综合分", "等级",
                    "红线预警", "主要风险点", "原始返回"])


def to_number(value):
    """把单元格值转成数字；转不了就原样返回字符串。"""
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return value
    s = str(value).strip().replace(",", "")
    try:
        return float(s) if "." in s else int(s)
    except ValueError:
        return value


def to_str_number(value):
    """把接口返回的数字安全转成字符串存表；None/空 转为空字符串。"""
    if value is None or value == "":
        return ""
    return str(value)


def normalize_level(level):
    """等级归一化：'高风险'/'高' → '高'，'中风险'/'中' → '中'，'低风险'/'低' → '低'。"""
    if level is None:
        return ""
    s = str(level).strip()
    if not s or s in ("-", "—"):
        return ""
    for word in ("高", "中", "低"):
        if word in s:
            return word
    return s  # 兜底：原样返回


def extract_redline(report_text):
    """从报告文本中提取'红线预警'冒号后的内容（到行尾为止），提取不到留空。"""
    if not report_text:
        return ""
    m = re.search(r"红线预警[：:]?\s*([^\n]+)", report_text)
    return m.group(1).strip() if m else ""


def call_api(payload):
    """调用接口，带重试。成功返回(resp_json, raw_text)，失败抛最后一次异常。"""
    last_err = None
    for attempt in range(MAX_RETRY + 1):
        try:
            resp = requests.post(API_URL, json=payload, timeout=REQUEST_TIMEOUT,
                                 headers={"Authorization": f"Bearer {API_TOKEN}"})
            if resp.status_code != 200:
                print(f"    [服务器返回 {resp.status_code}] {resp.text[:300]}")
            resp.raise_for_status()
            return resp.json(), resp.text
        except Exception as e:
            last_err = e
            if attempt < MAX_RETRY:
                wait = ROW_INTERVAL * (attempt + 1)
                print(f"    [警告] 请求失败（{e}），{wait}秒后第{attempt + 2}次尝试...")
                time.sleep(wait)
    raise last_err


def load_input_rows():
    """读取'案例字段'sheet，返回 [(case_id, payload), ...]。"""
    wb_in = load_workbook(INPUT_FILE, data_only=True)
    ws_in = wb_in[INPUT_SHEET]
    rows = []
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None for v in row):
            continue  # 跳过空行
        case_id = str(row[0]).strip() if row[0] is not None else ""
        payload = {}
        for i, key in enumerate(FIELD_KEYS, start=1):
            value = row[i] if i < len(row) else None
            payload[key] = to_number(value) if key in NUMERIC_FIELDS else ("" if value is None else str(value))
        rows.append((case_id, payload))
    return rows, wb_in


def load_existing_results():
    """
    中断重跑的关键：如果输出文件已存在，读取其中已完成的案例。
    返回 (results列表, 已完成case_id集合)。请求失败的行不算完成，重跑会重新调用。
    """
    results = []
    done_ids = set()
    if not __import__("os").path.exists(OUTPUT_FILE):
        return results, done_ids
    try:
        wb = load_workbook(OUTPUT_FILE, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or row[0] is None:
                continue
            case_id = str(row[0]).strip()
            raw = str(row[-1] or "")
            if raw.startswith("[请求失败]"):
                continue  # 失败行不算完成，重跑重试
            results.append(list(row[:len(OUTPUT_HEADER)]))
            done_ids.add(case_id)
    except Exception as e:
        print(f"[提示] 读取已有结果文件失败（{e}），将从头跑。")
        return [], set()
    return results, done_ids


def save_results(results):
    """把当前全部结果写入输出文件（每完成一条就保存一次，保证中断不丢数据）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "评测结果"
    ws.append(OUTPUT_HEADER)
    for r in results:
        ws.append(r)
    wb.save(OUTPUT_FILE)


def main():
    rows, wb_in = load_input_rows()
    total = len(rows)
    print(f"共读取到 {total} 条案例。")

    # 加载已有进度，跳过已完成的案例
    results, done_ids = load_existing_results()
    if done_ids:
        print(f"检测到 {len(done_ids)} 条已完成案例，本次将跳过（中断重跑模式）。")
    todo = [(cid, p) for cid, p in rows if cid not in done_ids]
    print(f"本次需调用 {len(todo)} 条，开始...\n")

    # ---------- 逐行调用并提取 ----------
    for n, (case_id, payload) in enumerate(todo, start=1):
        try:
            resp_json, raw_text = call_api(payload)
            report_text = str(resp_json.get("risk_report", "") or "")
            # 直读结构化字段，正则只用于红线预警
            out_row = (
                [case_id] + [payload[k] for k in FIELD_KEYS]
                + [to_str_number(resp_json.get("project_score")),   # 项目分
                   to_str_number(resp_json.get("credit_score")),    # 企业分
                   to_str_number(resp_json.get("contract_score")),  # 合同分
                   to_str_number(resp_json.get("total_score")),     # 综合分
                   normalize_level(resp_json.get("credit_advice")), # 等级
                   extract_redline(report_text),                    # 红线预警
                   str(resp_json.get("major_risk_points", "") or ""),  # 主要风险点
                   raw_text]
            )
            print(f"第{n}/{len(todo)}条 {case_id} 完成  "
                  f"综合分={out_row[16] or '空'} 等级={out_row[17] or '空'}")
        except Exception as e:
            # 失败不中断：该行标记"请求失败"，重跑时会自动重试
            out_row = ([case_id] + [payload[k] for k in FIELD_KEYS]
                       + ["", "", "", "", "", "", "", f"[请求失败] {type(e).__name__}: {e}"])
            print(f"第{n}/{len(todo)}条 {case_id} 请求失败（已重试{MAX_RETRY}次）: {e}")

        results.append(out_row)
        save_results(results)  # 每完成一条立即落盘，中断不丢进度
        if n < len(todo):
            time.sleep(ROW_INTERVAL)

    print(f"\n结果已保存到 {OUTPUT_FILE}\n")

    # ---------- 与专家评分对比 ----------
    if SUMMARY_SHEET not in wb_in.sheetnames:
        print(f"[警告] 未找到sheet'{SUMMARY_SHEET}'，跳过对比环节。")
        return

    ws_sum = wb_in[SUMMARY_SHEET]
    headers = [str(c.value).strip() if c.value else "" for c in ws_sum[1]]

    def find_col(*keywords):
        """按关键词模糊找列，兼容'专家-项目分'、'专家项目分'等写法。"""
        for i, h in enumerate(headers):
            if all(k in h for k in keywords):
                return i
        return None

    col_map = {
        "exp_proj": find_col("专家", "项目分"),
        "exp_ent": find_col("专家", "企业分"),
        "exp_contract": find_col("专家", "合同分"),
        "exp_total": find_col("专家", "综合分"),
        "exp_level": find_col("专家", "等级"),
    }
    try:
        id_col = headers.index("案例ID")
    except ValueError:
        print("[警告] 评测总表中没有'案例ID'列，跳过对比。")
        return
    for name, c in col_map.items():
        if c is None:
            print(f"[警告] 未找到列：{name}（表头：{headers}），该项指标跳过。")

    expert = {}  # case_id -> 整行
    for row in ws_sum.iter_rows(min_row=2, values_only=True):
        if row[id_col] is None:
            continue
        expert[str(row[id_col]).strip()] = row

    # 逐条匹配统计
    level_correct = level_matched = 0
    high_total = high_missed = 0
    mae_sum = mae_count = 0.0
    # 三个节点分的一致率统计：key -> [一致数, 可比数]
    node_stat = {"项目分": [0, 0], "企业分": [0, 0], "合同分": [0, 0]}
    node_col = {"项目分": col_map["exp_proj"], "企业分": col_map["exp_ent"], "合同分": col_map["exp_contract"]}
    agent_col = {"项目分": 13, "企业分": 14, "合同分": 15}  # 在输出列中的下标
    high_missed_list = []   # 高危漏报案例
    level_diff_list = []    # 等级分歧案例
    unmatched = 0

    for r in results:
        case_id = str(r[0]).strip()
        if case_id not in expert:
            print(f"[警告] 案例ID '{case_id}' 在评测总表中未找到，跳过对比。")
            unmatched += 1
            continue
        erow = expert[case_id]

        exp_level = normalize_level(erow[col_map["exp_level"]] if col_map["exp_level"] is not None else "")
        agent_level = str(r[17] or "")

        # 指标1：等级准确率
        if agent_level:
            level_matched += 1
            if agent_level == exp_level:
                level_correct += 1
            else:
                level_diff_list.append((case_id, exp_level, agent_level))

        # 指标2：高危漏报
        if exp_level == "高":
            high_total += 1
            if agent_level != "高":
                high_missed += 1
                high_missed_list.append((case_id, exp_level, agent_level or "未提取"))

        # 指标3：综合分MAE
        if col_map["exp_total"] is not None:
            try:
                mae_sum += abs(float(r[16]) - float(erow[col_map["exp_total"]]))
                mae_count += 1
            except (TypeError, ValueError):
                pass

        # 指标4：三个节点分一致率（|差|≤10）
        for node, ec in node_col.items():
            if ec is None:
                continue
            try:
                diff = abs(float(r[agent_col[node]]) - float(erow[ec]))
            except (TypeError, ValueError):
                continue  # 任一侧不是数字，不计入
            node_stat[node][1] += 1
            if diff <= 10:
                node_stat[node][0] += 1

    # ---------- 打印对比指标 ----------
    print("========== 对比指标 ==========")
    if unmatched:
        print(f"有 {unmatched} 条案例ID未能匹配到评测总表（详见上方警告）。")
    if level_matched:
        print(f"1. 综合等级判定准确率：{level_correct}/{level_matched} = {level_correct / level_matched:.1%}")
    else:
        print("1. 综合等级判定准确率：无有效数据。")
    if high_total:
        print(f"2. 高危漏报率：{high_missed}/{high_total} = {high_missed / high_total:.1%}")
    else:
        print("2. 高危漏报率：专家判'高'的案例数为0，不适用。")
    if mae_count:
        print(f"3. 综合分平均绝对误差（MAE）：{mae_sum / mae_count:.2f} 分（样本数 {int(mae_count)}）")
    else:
        print("3. 综合分平均绝对误差：无有效数据。")
    for node, (ok, cnt) in node_stat.items():
        if cnt:
            print(f"4. {node}一致率（|差|≤10）：{ok}/{cnt} = {ok / cnt:.1%}")
        else:
            print(f"4. {node}一致率：无可比数据。")

    # 高危漏报案例清单
    if high_missed_list:
        print("\n----- 高危漏报案例清单 -----")
        for cid, exp_lv, agent_lv in high_missed_list:
            print(f"  {cid}  专家={exp_lv}  Agent={agent_lv}")
    # 等级分歧案例清单
    if level_diff_list:
        print("\n----- 等级分歧案例清单 -----")
        for cid, exp_lv, agent_lv in level_diff_list:
            print(f"  {cid}  专家={exp_lv}  Agent={agent_lv}")
    print("================================")


if __name__ == "__main__":
    main()